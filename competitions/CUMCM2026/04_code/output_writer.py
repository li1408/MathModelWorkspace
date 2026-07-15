"""D 题结果文件写出工具。"""

from __future__ import annotations

import shutil
from copy import copy
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell import Cell

from data_io import MineNetwork, natural_edge_sort, natural_node_sort
from escape_routing import EscapeRoute
from water_flow import WaterResult


PROJECT_ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = PROJECT_ROOT / "02_raw_data" / "attachment_3"


def round_or_na(value: float | None) -> float | None:
    """按题目要求保留两位小数。"""
    if value is None:
        return None
    return round(float(value), 2)


def write_water_result_workbook(path: Path, network: MineNetwork, result: WaterResult) -> None:
    """复制官方模板并回填 result1/result3 类型的水流结果。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    _copy_template(path)
    workbook = load_workbook(path)
    node_sheet = workbook["端点水流到达时刻"]
    for row_index, node_id in enumerate(natural_node_sort(network.nodes), start=2):
        node_sheet.cell(row_index, 1).value = node_id
        node_sheet.cell(row_index, 2).value = round_or_na(result.node_arrival.get(node_id))
        node_sheet.cell(row_index, 2).number_format = "0.00"
    edge_sheet = workbook["巷道水流充满时刻"]
    for row_index, edge in enumerate(natural_edge_sort(network.edges), start=2):
        edge_sheet.cell(row_index, 1).value = edge.edge_id
        edge_sheet.cell(row_index, 2).value = round_or_na(result.edge_fill.get(edge.edge_id))
        edge_sheet.cell(row_index, 2).number_format = "0.00"
    workbook.save(path)


def write_escape_result_workbook(path: Path, routes: list[EscapeRoute]) -> None:
    """复制官方模板并回填 result2/result4 类型的完整逃生路径。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    _copy_template(path)
    workbook = load_workbook(path)
    for index, route in enumerate(routes, start=1):
        sheet = workbook[f"工人{index}逃生路径"]
        data_style = [_cell_style(sheet.cell(3, column)) for column in range(1, 7)]
        final_style = [_cell_style(sheet.cell(sheet.max_row, column)) for column in range(1, 7)]
        data_row_height = sheet.row_dimensions[3].height
        final_row_height = sheet.row_dimensions[sheet.max_row].height
        sheet.delete_rows(3, sheet.max_row - 2)

        start_point = route.start_point or (
            route.steps[0].x,
            route.steps[0].y,
            route.steps[0].z,
        )
        rows: list[list[object | None]] = [
            [0, start_point[0], start_point[1], start_point[2], route.notice_time, None]
        ]
        for step_index, step in enumerate(route.steps[1:-1], start=1):
            rows.append(
                [
                    step_index,
                    step.x,
                    step.y,
                    step.z,
                    step.time_min,
                    step.base_edge_id,
                ]
            )
        exit_point = route.exit_point or (
            route.steps[-1].x,
            route.steps[-1].y,
            route.steps[-1].z,
        )
        rows.append(
            ["最后", exit_point[0], exit_point[1], exit_point[2], route.arrival_time, route.exit_index]
        )
        for row_offset, values in enumerate(rows, start=3):
            style = final_style if values[0] == "最后" else data_style
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row_offset, column)
                cell.value = round(float(value), 2) if column in (2, 3, 4, 5) else value
                _apply_cell_style(cell, style[column - 1])
                if column in (2, 3, 4, 5):
                    cell.number_format = "0.00"
            sheet.row_dimensions[row_offset].height = (
                final_row_height if values[0] == "最后" else data_row_height
            )
    workbook.save(path)


def _copy_template(path: Path) -> None:
    template = TEMPLATE_DIR / path.name
    if not template.exists():
        raise FileNotFoundError(f"missing official result template: {template}")
    shutil.copy2(template, path)


def _cell_style(cell: Cell) -> dict[str, object]:
    return {
        "style": copy(cell._style),
        "number_format": cell.number_format,
        "alignment": copy(cell.alignment),
    }


def _apply_cell_style(cell: Cell, style: dict[str, object]) -> None:
    cell._style = copy(style["style"])
    cell.number_format = str(style["number_format"])
    cell.alignment = copy(style["alignment"])


def write_water_csvs(output_dir: Path, network: MineNetwork, result: WaterResult, prefix: str) -> None:
    """同时输出便于校验的 CSV 文件。"""
    output_dir.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(
        [
            {"node_id": node_id, "arrival_min": result.node_arrival.get(node_id)}
            for node_id in natural_node_sort(network.nodes)
        ]
    ).to_csv(output_dir / f"{prefix}_node_arrival.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(
        [
            {
                "edge_id": edge.edge_id,
                "entry_min": result.edge_entry.get(edge.edge_id),
                "fill_min": result.edge_fill.get(edge.edge_id),
            }
            for edge in natural_edge_sort(network.edges)
        ]
    ).to_csv(output_dir / f"{prefix}_edge_water.csv", index=False, encoding="utf-8-sig")


def write_route_csv(output_dir: Path, route: EscapeRoute, prefix: str) -> None:
    """写出单条逃生路径 CSV。"""
    output_dir.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(
        [
            {
                "worker": route.worker_name,
                "step": index,
                "node_id": step.node_id,
                "x": step.x,
                "y": step.y,
                "z": step.z,
                "time_min": step.time_min,
                "edge_id": step.base_edge_id,
                "internal_edge_id": step.edge_id,
                "reached_exit": route.reached_exit,
            }
            for index, step in enumerate(route.steps)
        ]
    ).to_csv(output_dir / f"{prefix}_{route.worker_name}.csv", index=False, encoding="utf-8-sig")
