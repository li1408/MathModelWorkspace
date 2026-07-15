"""Validate the manually maintained claim-to-evidence tables."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

CODE_DIR = Path(__file__).resolve().parents[1]
if str(CODE_DIR) not in sys.path:
    sys.path.insert(0, str(CODE_DIR))

from quality_gates import GateReport
from run_context import write_json


CLAIM_FIELDS = {
    "claim_id",
    "paper_section",
    "claim_text",
    "importance",
    "reviewer",
    "status",
}
EVIDENCE_FIELDS = {
    "evidence_id",
    "claim_id",
    "source_run_id",
    "source_result",
    "metric",
    "figure_or_table",
    "file_hash",
    "validation_method",
    "validation_status",
    "reviewer",
    "status",
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_csv(path: Path, required_fields: set[str]) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fields = set(reader.fieldnames or [])
        missing = required_fields - fields
        if missing:
            raise ValueError(f"{path.name} is missing columns: {sorted(missing)}")
        return list(reader)


def metric_is_readable(path: Path, metric: str) -> bool | None:
    if not metric:
        return None
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return metric in set(csv.DictReader(handle).fieldnames or [])
    if suffix == ".json":
        value: Any = json.loads(path.read_text(encoding="utf-8"))
        try:
            for part in metric.split("."):
                value = value[part]
            return True
        except (KeyError, TypeError):
            return False
    if suffix in {".xlsx", ".xlsm"}:
        if "!" not in metric:
            return None
        sheet_name, column = metric.split("!", 1)
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                return False
            headers = [cell.value for cell in next(workbook[sheet_name].iter_rows(max_row=1))]
            return column in headers
        finally:
            workbook.close()
    return None


def validate_evidence_tables(
    project_root: Path,
    claims_path: Path,
    links_path: Path,
    *,
    profile: str,
) -> GateReport:
    report = GateReport(stage="evidence_map")
    try:
        claims = read_csv(claims_path, CLAIM_FIELDS)
        links = read_csv(links_path, EVIDENCE_FIELDS)
    except (OSError, csv.Error, ValueError) as exc:
        report.add("ERROR", "evidence_table_invalid", str(exc))
        return report

    claim_ids: set[str] = set()
    for claim in claims:
        claim_id = claim["claim_id"].strip()
        if not claim_id or claim_id in claim_ids:
            report.add(
                "ERROR",
                "claim_id_invalid",
                "Claim IDs must be non-empty and unique.",
                claim_id=claim_id,
            )
        claim_ids.add(claim_id)
        if not claim["reviewer"].strip() or not claim["status"].strip():
            report.add(
                "WARNING",
                "claim_review_incomplete",
                "Claim reviewer or status is empty.",
                claim_id=claim_id,
            )

    verified_claims: set[str] = set()
    evidence_ids: set[str] = set()
    for link in links:
        evidence_id = link["evidence_id"].strip()
        claim_id = link["claim_id"].strip()
        source_run_id = link["source_run_id"].strip()
        source_result = link["source_result"].strip().replace("\\", "/")
        if not evidence_id or evidence_id in evidence_ids:
            report.add(
                "ERROR",
                "evidence_id_invalid",
                "Evidence IDs must be non-empty and unique.",
                evidence_id=evidence_id,
            )
        evidence_ids.add(evidence_id)
        if claim_id not in claim_ids:
            report.add(
                "ERROR",
                "evidence_claim_missing",
                "Evidence references an unknown claim.",
                evidence_id=evidence_id,
                claim_id=claim_id,
            )
        expected_prefix = f"05_model_results/runs/{source_run_id}/"
        if not source_run_id or not source_result.startswith(expected_prefix):
            report.add(
                "ERROR",
                "evidence_source_run_mismatch",
                "Evidence source_result must belong to its source_run_id directory.",
                evidence_id=evidence_id,
            )
            continue
        target = project_root / source_result
        if not target.is_file():
            report.add(
                "ERROR",
                "evidence_result_missing",
                "Registered evidence result file does not exist.",
                path=source_result,
                evidence_id=evidence_id,
            )
            continue
        actual_hash = sha256_file(target)
        if not link["file_hash"].strip() or actual_hash.lower() != link["file_hash"].strip().lower():
            report.add(
                "ERROR",
                "evidence_hash_mismatch",
                "Evidence SHA256 does not match the registered file hash.",
                path=source_result,
                evidence_id=evidence_id,
            )
            continue
        readable = metric_is_readable(target, link["metric"].strip())
        if readable is False:
            report.add(
                "ERROR",
                "evidence_metric_missing",
                "Registered metric cannot be read from the result file.",
                path=source_result,
                evidence_id=evidence_id,
                metric=link["metric"],
            )
            continue
        if readable is None and link["metric"].strip():
            report.add(
                "WARNING",
                "evidence_metric_manual_review",
                "The metric format requires manual verification.",
                path=source_result,
                evidence_id=evidence_id,
            )
        if link["validation_status"].strip() == "verified":
            verified_claims.add(claim_id)
        if not link["reviewer"].strip() or not link["status"].strip():
            report.add(
                "WARNING",
                "evidence_review_incomplete",
                "Evidence reviewer or status is empty.",
                evidence_id=evidence_id,
            )

    severity = "ERROR" if profile == "final" else "WARNING"
    for claim in claims:
        if claim["importance"].strip() == "major" and claim["claim_id"].strip() not in verified_claims:
            report.add(
                severity,
                "major_claim_without_verified_evidence",
                "A major claim has no verified evidence link.",
                claim_id=claim["claim_id"],
            )
    report.metrics = {
        "claim_count": len(claims),
        "evidence_link_count": len(links),
        "verified_claim_count": len(verified_claims),
    }
    return report


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", default=os.environ.get("CUMCM_PROJECT_ROOT", "."))
    parser.add_argument("--run-dir", default=os.environ.get("CUMCM_RUN_DIR"))
    parser.add_argument("--profile", default=os.environ.get("CUMCM_PROFILE", "audit"))
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    project_root = Path(args.project_root).resolve()
    report = validate_evidence_tables(
        project_root,
        project_root / "07_paper/evidence/claims.csv",
        project_root / "07_paper/evidence/evidence_links.csv",
        profile=args.profile,
    )
    if args.run_dir:
        write_json(Path(args.run_dir) / "evidence" / "evidence_map_report.json", report.to_dict())
    return report.exit_code()


if __name__ == "__main__":
    raise SystemExit(main())
