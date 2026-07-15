"""Run smoke and full-local cold reproductions in fresh E-drive workspaces."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import uuid
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

CODE_DIR = Path(__file__).resolve().parents[1]
if str(CODE_DIR) not in sys.path:
    sys.path.insert(0, str(CODE_DIR))

from quality_gates import GateIssue
from run_context import sha256_file, write_json


SMOKE_RUNNER = r'''from __future__ import annotations
import json
from pathlib import Path
from graph_utils import GraphEdge, GraphModel, GraphNode
from water_flow import SourceSpec, simulate_water
from escape_routing import find_escape_route

nodes = {
    "S": GraphNode("S", 0.0, 0.0, 0.0),
    "T": GraphNode("T", 100.0, 0.0, 0.0),
}
graph = GraphModel(0, nodes, {"E0": GraphEdge("E0", "S", "T", 100.0, "E0")})
water = simulate_water(graph, [SourceSpec("source", nodes["S"].point, 0.0)])
route = find_escape_route(
    graph,
    water,
    worker_name="fixture_worker",
    start_point=(50.0, 0.0, 0.0),
    exit_points=[nodes["T"].point],
    notice_time=0.0,
)
result = {
    "edge_fill_min": water.edge_fill["E0"],
    "mass_balance_error_m3": water.max_mass_balance_error_m3,
    "route_reached_exit": route.reached_exit,
    "route_exit_index": route.exit_index,
    "route_arrival_min": route.arrival_time,
}
Path("result.json").write_text(
    json.dumps(result, sort_keys=True, separators=(",", ":")) + "\n",
    encoding="utf-8",
)
'''

SMOKE_MODULES = (
    "data_io.py",
    "graph_utils.py",
    "flow_strategies.py",
    "water_flow.py",
    "escape_routing.py",
)


def run_smoke_once(project_root: Path, workspace: Path) -> Path:
    workspace.mkdir(parents=True, exist_ok=False)
    for name in SMOKE_MODULES:
        shutil.copy2(project_root / "04_code" / name, workspace / name)
    (workspace / "smoke_runner.py").write_text(SMOKE_RUNNER, encoding="utf-8")
    env = os.environ.copy()
    env.pop("PYTHONPATH", None)
    for key in list(env):
        if key.startswith("CUMCM_"):
            env.pop(key, None)
    subprocess.run(
        [sys.executable, "smoke_runner.py"],
        cwd=workspace,
        env=env,
        check=True,
        capture_output=True,
        text=True,
    )
    result = workspace / "result.json"
    if not result.is_file():
        raise RuntimeError("Smoke fixture did not produce result.json.")
    return result


def run_smoke_fixture(project_root: Path, output_root: Path) -> dict[str, Any]:
    token = uuid.uuid4().hex[:10]
    first = run_smoke_once(project_root, output_root / f"smoke_{token}_a")
    second = run_smoke_once(project_root, output_root / f"smoke_{token}_b")
    first_hash = sha256_file(first)
    second_hash = sha256_file(second)
    return {
        "mode": "smoke_fixture",
        "status": "PASS" if first_hash == second_hash else "FAIL",
        "first_sha256": first_hash,
        "second_sha256": second_hash,
        "first_result_json": first.read_text(encoding="utf-8"),
        "comparison": "CSV/JSON-style exact SHA256 comparison",
    }


def workbook_content(path: Path) -> dict[str, list[list[Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        return {
            sheet_name: [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
            for sheet_name in workbook.sheetnames
        }
    finally:
        workbook.close()


def compare_excel_workbooks(expected: Path, actual: Path) -> dict[str, Any]:
    expected_content = workbook_content(expected)
    actual_content = workbook_content(actual)
    return {
        "equal": expected_content == actual_content,
        "expected_sheets": list(expected_content),
        "actual_sheets": list(actual_content),
        "comparison": "workbook sheet names and cell content; file bytes are not compared",
    }


def resolve_input_root() -> Path:
    value = os.environ.get("CUMCM_INPUT_ROOT")
    if not value:
        raise ValueError("CUMCM_INPUT_ROOT is required for full_local reproduction.")
    root = Path(value).expanduser().resolve()
    if not root.is_dir():
        raise ValueError("CUMCM_INPUT_ROOT does not point to a directory.")
    candidate = root / "02_raw_data"
    return candidate if candidate.is_dir() else root


def resolve_latexmk(env: dict[str, str]) -> str | None:
    """Prefer the MiKTeX runtime injected by RunContext over the system PATH."""
    miktex_bin = env.get("MIKTEX_BIN", "").strip()
    if miktex_bin:
        candidate = Path(miktex_bin) / "latexmk.exe"
        if candidate.is_file():
            return str(candidate)
    return shutil.which("latexmk", path=env.get("PATH"))


def _write_process_log(path: Path, process: subprocess.CompletedProcess[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        "[stdout]\n"
        + (process.stdout or "")
        + "\n[stderr]\n"
        + (process.stderr or ""),
        encoding="utf-8",
    )


def _copy_tree(source: Path, target: Path, *, ignore: tuple[str, ...] = ()) -> None:
    shutil.copytree(
        source,
        target,
        ignore=shutil.ignore_patterns("__pycache__", "*.pyc", *ignore),
    )


def compare_pdf_semantics(expected: Path, actual: Path) -> dict[str, Any]:
    pdfinfo = shutil.which("pdfinfo")
    pdftotext = shutil.which("pdftotext")
    if not pdfinfo or not pdftotext:
        return {"equal": False, "status": "TOOLS_MISSING", "comparison": "pages and extracted text"}

    def inspect(path: Path) -> dict[str, Any]:
        info = subprocess.run(
            [pdfinfo, str(path)], check=True, capture_output=True, text=True, errors="replace"
        ).stdout
        page_match = re.search(r"^Pages:\s+(\d+)", info, flags=re.MULTILINE)
        text = subprocess.run(
            [pdftotext, str(path), "-"], check=True, capture_output=True
        ).stdout
        return {
            "pages": int(page_match.group(1)) if page_match else None,
            "text_sha256": hashlib.sha256(text).hexdigest(),
        }

    expected_info = inspect(expected)
    actual_info = inspect(actual)
    return {
        "equal": expected_info == actual_info,
        "status": "CHECKED",
        "expected": expected_info,
        "actual": actual_info,
        "comparison": "page count and extracted text; PDF bytes are not compared",
    }


def run_full_local(project_root: Path, output_root: Path) -> dict[str, Any]:
    input_root = resolve_input_root()
    workspace = output_root / f"full_local_{uuid.uuid4().hex[:10]}"
    workspace.mkdir(parents=True, exist_ok=False)
    _copy_tree(input_root, workspace / "02_raw_data")
    for name in ("04_code", "config", "01_problem", "styles", "06_paper_assets", "07_paper"):
        source = project_root / name
        if source.is_dir():
            ignore = ("build",) if name == "07_paper" else ()
            _copy_tree(source, workspace / name, ignore=ignore)
    for name in ("requirements.txt", ".gitignore"):
        source = project_root / name
        if source.is_file():
            shutil.copy2(source, workspace / name)

    env = os.environ.copy()
    env.pop("PYTHONPATH", None)
    for key in list(env):
        if key.startswith("CUMCM_"):
            env.pop(key, None)
    model = subprocess.run(
        [sys.executable, "04_code/run_all.py"],
        cwd=workspace,
        env=env,
        capture_output=True,
        text=True,
    )
    model_log = output_root / "full_local_model.log"
    _write_process_log(model_log, model)
    comparisons: list[dict[str, Any]] = []
    for expected in sorted((project_root / "05_model_results/metrics").glob("*.csv")):
        actual = workspace / "05_model_results/metrics" / expected.name
        comparisons.append(
            {
                "file": f"05_model_results/metrics/{expected.name}",
                "kind": "csv_exact_sha256",
                "equal": actual.is_file() and sha256_file(expected) == sha256_file(actual),
            }
        )
    for expected in sorted((project_root / "05_model_results/tables").glob("result*.xlsx")):
        actual = workspace / "05_model_results/tables" / expected.name
        comparison = (
            compare_excel_workbooks(expected, actual)
            if actual.is_file()
            else {"equal": False, "comparison": "actual workbook missing"}
        )
        comparisons.append(
            {"file": f"05_model_results/tables/{expected.name}", "kind": "excel_content", **comparison}
        )

    latexmk = resolve_latexmk(env)
    paper_compile: dict[str, Any] = {"status": "NOT_RUN", "reason": "latexmk not found on PATH"}
    if latexmk:
        compile_process = subprocess.run(
            [latexmk, "-xelatex", "-outdir=build", "main.tex"],
            cwd=workspace / "07_paper",
            env=env,
            capture_output=True,
            text=True,
        )
        latex_log = output_root / "full_local_latexmk.log"
        _write_process_log(latex_log, compile_process)
        actual_pdf = workspace / "07_paper/build/main.pdf"
        expected_pdf = project_root / "07_paper/build/main.pdf"
        paper_compile = {
            "status": "PASS" if compile_process.returncode == 0 and actual_pdf.is_file() else "FAIL",
            "return_code": compile_process.returncode,
            "log": latex_log.relative_to(output_root.parent).as_posix(),
        }
        if actual_pdf.is_file() and expected_pdf.is_file():
            paper_compile["semantic_comparison"] = compare_pdf_semantics(expected_pdf, actual_pdf)

    passed = (
        model.returncode == 0
        and all(bool(item.get("equal")) for item in comparisons)
        and paper_compile.get("status") == "PASS"
        and bool(paper_compile.get("semantic_comparison", {}).get("equal"))
    )
    return {
        "mode": "full_local",
        "status": "PASS" if passed else "FAIL",
        "workspace": workspace.relative_to(output_root.parent).as_posix(),
        "model_return_code": model.returncode,
        "model_log": model_log.relative_to(output_root.parent).as_posix(),
        "comparisons": comparisons,
        "paper_compile": paper_compile,
    }


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", default=os.environ.get("CUMCM_PROJECT_ROOT"))
    parser.add_argument("--run-dir", default=os.environ.get("CUMCM_RUN_DIR"))
    parser.add_argument("--profile", default=os.environ.get("CUMCM_PROFILE", "audit"))
    parser.add_argument("--mode", choices=("smoke_fixture", "full_local", "both"), default=None)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    if not args.project_root or not args.run_dir:
        print("[CONFIG ERROR] CUMCM_PROJECT_ROOT and CUMCM_RUN_DIR are required.", file=sys.stderr)
        return 2
    project_root = Path(args.project_root).resolve()
    run_dir = Path(args.run_dir).resolve()
    output_root = run_dir / "cold_reproduction"
    mode = args.mode or ("both" if args.profile == "final" else "smoke_fixture")
    reports: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    try:
        if mode in {"smoke_fixture", "both"}:
            smoke = run_smoke_fixture(project_root, output_root)
            reports.append(smoke)
            if smoke["status"] != "PASS":
                issues.append(
                    GateIssue("ERROR", "cold_smoke_failed", "Smoke fixture reproduction failed.").to_dict()
                )
        if mode in {"full_local", "both"}:
            full = run_full_local(project_root, output_root)
            reports.append(full)
            if full["status"] != "PASS":
                issues.append(
                    GateIssue("ERROR", "cold_full_local_failed", "Full-local reproduction failed.").to_dict()
                )
        write_json(
            output_root / "cold_reproduction_report.json",
            {
                "source_run_id": run_dir.name,
                "temporary_storage_policy": "All reproduction workspaces are under the E-drive run directory.",
                "reports": reports,
                "issues": issues,
            },
        )
        return 1 if issues else 0
    except ValueError as exc:
        write_json(
            output_root / "cold_reproduction_report.json",
            {
                "source_run_id": run_dir.name,
                "reports": reports,
                "issues": [
                    GateIssue("ERROR", "cold_reproduction_config", str(exc)).to_dict()
                ],
            },
        )
        return 1
    except (OSError, subprocess.SubprocessError) as exc:
        print(f"[SCRIPT ERROR] {type(exc).__name__}: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
