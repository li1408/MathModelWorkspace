from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


CODE_DIR = Path(__file__).resolve().parents[1]
PROJECT_ROOT = CODE_DIR.parent
if str(CODE_DIR) not in sys.path:
    sys.path.insert(0, str(CODE_DIR))


from d_workflow import (
    combine_adjusted_route,
    compare_time_monotonicity,
    replay_route_prefix,
    sha256_file,
)
from escape_routing import EscapeRoute, RouteStep
from graph_utils import GraphEdge, GraphModel, GraphNode
from output_writer import write_escape_result_workbook


class OutputContractTests(unittest.TestCase):
    def test_sha256_file_matches_known_digest(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.txt"
            path.write_bytes(b"abc")

            self.assertEqual(
                sha256_file(path),
                "ba7816bf8f01cfea414140de5dae2223"
                "b00361a396177a9cb410ff61f20015ad",
            )

    def test_q3_monotonicity_check_rejects_later_times(self):
        errors = compare_time_monotonicity(
            {"A": 1.0, "B": 2.0, "C": None},
            {"A": 0.5, "B": 2.5, "C": 3.0},
            label="node",
        )

        self.assertEqual(errors, ["node B: q3=2.5 later than q1=2.0"])

    def test_q4_prefix_is_replayed_under_two_source_water_state(self):
        graph = GraphModel(
            mine_id=0,
            nodes={
                "A": GraphNode("A", 0.0, 0.0, 0.0),
                "B": GraphNode("B", 720.0, 0.0, 0.0),
            },
            edges={"H0001": GraphEdge("H0001", "A", "B", 720.0, "H0001")},
        )

        class Water:
            edge_entry = {"H0001": 2.0}

            @staticmethod
            def edge_safe_until(edge_id: str) -> float:
                return 30.0

            @staticmethod
            def edge_direction_sign(edge_id: str, time_min: float) -> int:
                return -1

            @staticmethod
            def edge_state_change_times(edge_id: str, after_time: float) -> list[float]:
                return [time for time in (2.0, 30.0) if time > after_time]

        original = EscapeRoute(
            worker_name="worker1",
            reached_exit=True,
            arrival_time=4.0,
            exit_node="B",
            exit_index=1,
            start_point=(0.0, 0.0, 0.0),
            exit_point=(720.0, 0.0, 0.0),
            notice_time=1.0,
            steps=[
                RouteStep("A", 0.0, 0.0, 0.0, 1.0, None, None),
                RouteStep("B", 720.0, 0.0, 0.0, 4.0, "H0001", "H0001"),
            ],
        )

        prefix = replay_route_prefix(original, graph, Water(), 1.0, 5.0)

        self.assertAlmostEqual(prefix[-1].x, 420.0, places=8)
        self.assertAlmostEqual(prefix[-1].time_min, 5.0, places=8)
        self.assertEqual(prefix[-1].base_edge_id, "H0001")

    def test_escape_workbook_preserves_template_and_outputs_only_official_ids(self):
        route = EscapeRoute(
            worker_name="worker1",
            reached_exit=True,
            arrival_time=3.5,
            exit_node="V_exit_1",
            exit_index=1,
            start_point=(5808.18, 5367.75, 10.0),
            exit_point=(3252.16, 3326.63, 10.0),
            notice_time=1.0,
            steps=[
                RouteStep("V_start", 5808.19, 5367.74, 10.0, 1.0, None, None),
                RouteStep("P0001", 5500.0, 5100.0, 10.0, 2.0, "H0001__V_start_a", "H0001"),
                RouteStep("V_exit_1", 3252.15, 3326.62, 10.0, 3.5, "H0002__V_exit_a", "H0002"),
            ],
        )
        routes = [route, route, route]

        temp_root = PROJECT_ROOT / "output" / "test_tmp"
        temp_root.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=temp_root) as tmp:
            output_path = Path(tmp) / "result2-1.xlsx"
            write_escape_result_workbook(output_path, routes)
            generated = load_workbook(output_path)
            template = load_workbook(PROJECT_ROOT / "02_raw_data" / "attachment_3" / "result2-1.xlsx")

        ws = generated["工人1逃生路径"]
        template_ws = template["工人1逃生路径"]
        self.assertEqual(
            {str(cell_range) for cell_range in ws.merged_cells.ranges},
            {str(cell_range) for cell_range in template_ws.merged_cells.ranges},
        )
        self.assertEqual([ws.cell(1, col).value for col in range(1, 7)], ["序号", "工人位置", None, None, "到达时刻", "巷道编号"])
        self.assertNotIn("Unnamed: 2", [cell.value for row in ws.iter_rows() for cell in row])
        self.assertEqual(ws["B3"].value, 5808.18)
        self.assertEqual(ws["C3"].value, 5367.75)
        self.assertEqual(ws["E3"].value, 1.0)
        self.assertEqual(ws["F4"].value, "H0001")
        self.assertEqual(ws["A5"].value, "最后")
        self.assertEqual(ws["B5"].value, 3252.16)
        self.assertEqual(ws["F5"].value, 1)
        for row in range(3, 6):
            for column in range(2, 6):
                self.assertEqual(ws.cell(row, column).number_format, "0.00")

    def test_q4_combination_keeps_q2_prefix_and_adjustment_point(self):
        original = EscapeRoute(
            worker_name="worker1",
            reached_exit=True,
            arrival_time=10.0,
            exit_node="exit_old",
            exit_index=1,
            start_point=(0.0, 0.0, 0.0),
            exit_point=(90.0, 0.0, 0.0),
            notice_time=1.0,
            steps=[
                RouteStep("start", 0.0, 0.0, 0.0, 1.0, None, None),
                RouteStep("n1", 20.0, 0.0, 0.0, 3.0, "E1", "H0001"),
                RouteStep("n2", 60.0, 0.0, 0.0, 7.0, "E2", "H0002"),
                RouteStep("exit_old", 90.0, 0.0, 0.0, 10.0, "E3", "H0003"),
            ],
        )
        adjusted = EscapeRoute(
            worker_name="worker1",
            reached_exit=True,
            arrival_time=8.0,
            exit_node="exit_new",
            exit_index=2,
            start_point=(40.0, 0.0, 0.0),
            exit_point=(50.0, 50.0, 0.0),
            notice_time=5.0,
            steps=[
                RouteStep("adjust", 40.0, 0.0, 0.0, 5.0, None, None),
                RouteStep("exit_new", 50.0, 50.0, 0.0, 8.0, "E4", "H0004"),
            ],
        )

        combined = combine_adjusted_route(original, adjusted, adjustment_time=5.0)

        self.assertEqual(combined.notice_time, 1.0)
        self.assertEqual(combined.start_point, original.start_point)
        self.assertEqual(combined.exit_point, adjusted.exit_point)
        self.assertEqual([step.time_min for step in combined.steps], [1.0, 3.0, 5.0, 8.0])
        self.assertEqual(combined.steps[2].base_edge_id, "H0002")


if __name__ == "__main__":
    unittest.main()
